import os
import io
import json
import base64
import zipfile
import tempfile

from pathlib import Path
from flask import Flask, request
from flask_cors import CORS

from gliner import GLiNER
from langdetect import detect

app = Flask(__name__)
CORS(app)

################################
# 1. LOAD GLiNER MODELS
################################
gliner_models = {
    "en": GLiNER.from_pretrained("gliner-community/gliner_medium-v2.5"),
    "it": GLiNER.from_pretrained("DeepMount00/GLiNER_ITA_LARGE")
}

################################
# FUNZIONE PER CARICARE LABELS
################################
def load_labels():
    """
    Carica dinamicamente le etichette (entities) da gliner_entities.txt.
    Se il file non esiste, ritorna un fallback.
    """
    if os.path.exists("gliner_entities.txt"):
        with open("gliner_entities.txt", "r", encoding="utf-8") as f:
            labels = [line.strip() for line in f if line.strip()]
        if labels:
            return labels
    # Fallback
    return ["PERSON", "ORG", "LOC"]

############################
# 2. UTILITY FUNCTIONS
############################

def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    text = ""

    if suffix == ".pdf":
        from PyPDF2 import PdfReader
        try:
            reader = PdfReader(str(path))
            for page in reader.pages:
                page_text = page.extract_text() or ""
                text += page_text + "\n"
        except Exception as e:
            print(f"Error reading PDF {path}: {e}")
            return ""
        return text.strip()

    elif suffix == ".docx":
        import docx
        try:
            doc = docx.Document(str(path))
            for para in doc.paragraphs:
                text += para.text + "\n"
        except Exception as e:
            print(f"Error reading DOCX {path}: {e}")
            return ""
        return text.strip()

    elif suffix == ".xlsx":
        import openpyxl
        try:
            wb = openpyxl.load_workbook(str(path), data_only=True)
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    row_text = [str(cell) for cell in row if cell is not None]
                    if row_text:
                        text += " ".join(row_text) + "\n"
        except Exception as e:
            print(f"Error reading XLSX {path}: {e}")
            return ""
        return text.strip()

    elif suffix == ".txt":
        try:
            return path.read_text(encoding="utf-8", errors="replace")
        except Exception as e:
            print(f"Error reading TXT {path}: {e}")
            return ""
    else:
        print(f"Unsupported file format: {suffix}")
        return ""

def detect_language_of_text(text: str) -> str:
    try:
        code = detect(text)
        if code not in gliner_models:
            return "en"  # fallback
        return code
    except:
        return "en"  # fallback

def anonymize_text(text: str, labels, lang_code: str):
    """
    Usa GLiNER con la lista di labels passata.
    Ritorna (anonymized_text, mapping_list).
    """
    model = gliner_models[lang_code]
    entities = model.predict_entities(text, labels, threshold=0.5)

    anonymized = text
    mapping_list = []
    for i, ent in enumerate(entities):
        placeholder = f"{{{{{ent['label']}_{i}}}}}"
        anonymized = anonymized.replace(ent["text"], placeholder)
        mapping_list.append((placeholder, ent["text"]))
    return anonymized, mapping_list

def deanonymize_text(anonymized_text: str, mapping_list: list) -> str:
    deanonymized = anonymized_text
    for placeholder, original in mapping_list:
        deanonymized = deanonymized.replace(placeholder, original)
    return deanonymized

############################
# 3. ENDPOINT /api/anonymize
############################
@app.route("/api/anonymize", methods=["POST"])
def api_anonymize():
    # 1. Ricarica le etichette
    labels = load_labels()
    print(f"Using labels: {labels}")

    # 2. Ricevi il file
    if "file" not in request.files:
        return {"error": "File not provided"}, 400
    uploaded_file = request.files["file"]
    if uploaded_file.filename == "":
        return {"error": "Empty filename"}, 400

    # 3. Salva e leggi il testo
    with tempfile.NamedTemporaryFile(delete=False, suffix=uploaded_file.filename) as tf:
        uploaded_file.save(tf.name)
        raw_text = extract_text(Path(tf.name))

    # 4. Rileva lingua e anonimizza
    lang_code = detect_language_of_text(raw_text)
    anonymized_text, mapping_list = anonymize_text(raw_text, labels, lang_code)

    # 5. Crea ZIP con anonymized.txt e mapping.json
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("anonymized.txt", anonymized_text)
        zf.writestr("mapping.json", json.dumps({
            "mapping": mapping_list,
            "language": lang_code
        }, ensure_ascii=False, indent=2))

    zip_buffer.seek(0)
    zip_data = zip_buffer.read()
    zip_base64 = base64.b64encode(zip_data).decode("utf-8")

    return {
        "anonymized_text": anonymized_text,
        "zip_base64": zip_base64
    }

##############################
# 4. ENDPOINT /api/deanonymize
##############################
@app.route("/api/deanonymize", methods=["POST"])
def api_deanonymize():
    # Verifica che textFile e mappingFile siano presenti
    if "textFile" not in request.files or "mappingFile" not in request.files:
        return {"error": "Missing textFile or mappingFile"}, 400

    # Leggi il testo anonimizzato
    anonymized_data = request.files["textFile"].read().decode("utf-8", errors="replace")

    # Leggi il file di mapping
    mapping_content = request.files["mappingFile"].read().decode("utf-8", errors="replace")

    try:
        mapping_json = json.loads(mapping_content)
    except:
        return {"error": "Invalid JSON in mappingFile"}, 400

    mapping_list = mapping_json.get("mapping", [])
    # Esegui la de-anonimizzazione
    deanonymized = deanonymize_text(anonymized_data, mapping_list)

    # Codifica il testo de-anonimizzato in base64 (per l'anteprima)
    result_b64 = base64.b64encode(deanonymized.encode("utf-8")).decode("utf-8")

    # 1) Creiamo anche un ZIP con il file "deanonymized.txt" 
    #    (così l'utente può scaricarlo subito)
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("deanonymized.txt", deanonymized)

    zip_buffer.seek(0)
    zip_data = zip_buffer.read()
    zip_base64 = base64.b64encode(zip_data).decode("utf-8")

    # Ritorniamo sia la stringa in base64 (deanonymized_base64) 
    # che lo ZIP da scaricare (zip_base64)
    return {
        "deanonymized_base64": result_b64,
        "zip_base64": zip_base64
    }

###################
# RUN APP
###################
if __name__ == "__main__":
    app.run(debug=True, port=5000)
